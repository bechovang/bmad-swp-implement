# -*- coding: utf-8 -*-
"""Split Sprint_Backlog.xlsx Business Rules -> Business Rules (rules only) + Note (tech/impl)."""
import shutil, math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = r"D:\FPT_Ky5\SWP\bmad swp\Sprint_Backlog.xlsx"
BAK = r"D:\FPT_Ky5\SWP\bmad swp\.working\Sprint_Backlog_backup.xlsx"

BR = {
1: '''• contracts/openapi.yaml là chân lý API duy nhất (AD-2): đóng băng đầu sprint, muốn đổi thì PR vào contract trước.
• Lỗi + list theo envelope chung (AD-8): lỗi = machine code + message + field errors; list = {items, page, pageSize, total}, mặc định 25 rows.''',
2: '''• Tiền = DECIMAL(15,0) — cấm float/double (AD-7).
• Thời gian lưu UTC.
• ActivityLog append-only: chỉ INSERT, không UPDATE/DELETE.''',
3: '''• Seed đủ dữ liệu lịch sử để chạy dashboard + 6 UJ end-to-end.''',
4: '''• FR-1: JWT Bearer, role trong claim, TTL 24h, không refresh token.
• Landing theo role: Customer → Browse Units; Staff → Task Board; Facility Manager → Facility Overview; Business Ops → Business Overview; SysAdmin → User Management.
• Đăng nhập sai → message chung, không tiết lộ field nào sai; truy cập URL chéo role → 403.
• Permission matrix enforce server-side ở từng endpoint (AD-5).
• Mỗi lần đăng nhập (thành/bại) ghi ActivityLog LOGIN/LOGIN_FAILED (FR-38).''',
5: '''• FR-2: form công khai chỉ tạo CUSTOMER — không có trường chọn role.
• Email trùng chặn inline; thành công tự đăng nhập → Browse Units.
• FR-3 (P2): quên mật khẩu trả lời generic "If an account exists…", không gửi email thật.''',
6: '''• FR-4: filter loại / kích thước / ngày bắt đầu / thời hạn; sort mặc định giá thấp → cao.
• Chỉ re-query khi bấm Search (không per keystroke).
• Availability = Reservation + Turnover Buffer tính tại thời điểm query (AD-4) — FE hiển thị nguyên văn, không tính lại.
• Unit Rented/Maintenance/Retired không xuất hiện.
• Card buffer "Available {date} · cleaning buffer" vẫn đặt được với ngày đó.''',
7: '''• FR-6: bảng giá = rent × duration + từng dòng phụ thu từ Rental Policy active + Deposit đánh dấu refundable.
• Mọi dòng tiền khớp tuyệt đối Booking Summary và Payment Modal.
• Tiền VND full precision, format 1.150.000 ₫, tabular numerals (NFR-1).''',
8: '''• FR-8/9: một modal dùng chung 4 touchpoint (Deposit, 100% rent, Extension fee, Extra fee).
• 3 phương thức: Card (validation inline) / MoMo (phone + OTP) / VNPay QR (QR + đếm ngược ~5 phút).
• Processing ~1.5–2s: modal không đóng/điều hướng được.
• Fail = "No money was taken" + Retry + Switch method, KHÔNG đổi trạng thái gì anywhere; QR hết hạn → về method select; sau 2 fail gợi ý đổi phương thức.
• Payment thành công sinh đúng 1 receipt xem vĩnh viễn.''',
9: '''• FR-5/7: khách duyệt lại toàn bộ dòng sẽ bị truy đòi trước khi trả cọc — không có dòng tiền nào xuất hiện ở payment mà vắng ở đây.
• Bấm Reserve → hệ thống re-check availability: unit bị chiếm giữa chừng → quay về lưới + toast gợi ý unit tương tự; không bao giờ booking trên ngày invalid.
• Reserve sinh Reservation PENDING_PAYMENT; cọc thành công → RESERVED (unit bị giữ từ đây).
• Contract được auto-draft từ đúng các điều kiện đã duyệt và ký tại check-in.''',
10: '''• FR-36: hết ngày nhận kho mà chưa check-in → EXPIRED: mất cọc (receipt "Deposit forfeited — no-show"), Unit → Available, Contract Draft/Printed → Closed, ghi Activity Log + notification khách.
• EXPIRED tính on-read idempotent (AD-4) — API luôn trả state đã suy diễn xong, FE không tự suy từ ngày.
• Không có cancel trong v1 (PRD §6).''',
11: '''• FR-10: sinh Contract Draft ngay khi Deposit thành công; mã CT-xxxx; khóa phiên bản Rental Policy.
• Contract read-only ở mọi nơi — không UI nào cho phép sửa, không ai soạn contract tay.
• Sửa bản Draft = re-draft: bản cũ chuyển Superseded, vẫn đọc được trong chuỗi (FR-13).''',
12: '''• FR-35: My Rentals = Reservations + Rentals + History dạng card + status badge + nút action kế tiếp.
• Là pointer surface: mọi con số/trạng thái khớp nguồn, không tự tính lại.
• Check-in Pass: mã BK- hiển thị lớn (mono) để staff validate tại quầy + hướng dẫn plain-English (mang ID ký hợp đồng, trả 100% tiền thuê tại quầy).
• Empty state có đúng 1 CTA trỏ flow đổ dữ liệu (NFR-8).''',
13: '''• FR-9/13/35: một rental = unit, dates, payment history + receipts vĩnh viễn, deposit status (held/settled).
• Contract chain: contract gốc + addenda (mã CT-, badge trạng thái, ảnh bản ký); bản Signed xem mãi; Superseded vẫn đọc được; addendum chưa ký có amber banner nêu deadline tại quầy.
• Là anchor của mọi action: Check-in, Extend, Checkout Request, New Support — không flow nào bắt đầu từ system menu.''',
14: '''• FR-21: cột To do / In progress / Done; 5 loại card Check-in / Checkout / Cleaning / Support / Contract-signature; tab lọc theo loại; số lượng từng cột ở header; card Done render muted.
• Card sinh theo ca trực (route theo Unit + shift).
• Kéo thả HOẶC keyboard Move — drag không phải con đường duy nhất (NFR-2).''',
15: '''• FR-11/14/15: nhập/mã reservation → validate tồn tại + Deposit đã trả; không hợp lệ chặn kèm lý do cụ thể (chưa cọc không check-in được).
• Thu 100% rent qua Payment Modal — Deposit và rent full là 2 giá trị tách bạch.
• Contract ritual: preview read-only → Print → chụp ảnh bản ký → Attach (ghi CONTRACT_SIGNED vào Activity Log); nút trao Access Code VÔ HIỆU đến khi có ảnh bản ký.
• Hoàn tất → Rental CHECKED_IN, Unit RENTED, task Done.''',
16: '''• FR-22: kéo Check-in/Checkout/Contract-signature vào Done khi thiếu closing step (payment/settlement/ảnh bản ký) → card bật lại + toast nêu đúng bước thiếu.
• Không có task Done mà nghiệp vụ chưa khép — enforce phía service (AD-4), FE hiển thị kết quả.''',
17: '''• FR-33/34: toast tức thời ~4s, hover pause, ≤1 action link.
• Bell badge đếm unread; Notification Center feed theo role (khách: booking/payment/support; staff: task assignment; manager: escalation + status writes), unread trước, deep link tới đối tượng, mark-all-read, unread persist giữa phiên.
• Sự kiện tiền phát CẢ toast lẫn bell.
• Không email — không bao giờ.''',
18: '''• FR-16: khách chọn ngày checkout mới → hệ thống kiểm Reservation kế tiếp; vùng xung đột đánh dấu trên picker; submit bị chặn kèm biên rõ ("Latest new checkout: Oct 18").
• Chỉ mở TRƯỚC EndDate — qua EndDate giữ CHECKED_IN, mỗi ngày trễ tính phụ thu LATE_FEE tại Settlement (không có trạng thái OVERDUE).
• Phí gia hạn = tiền thuê kỳ thêm theo Rental Policy active (không dùng công thức cọc mới − cọc cũ); Deposit giữ nguyên booking gốc.
• Thanh toán thành công → EndDate flip NGAY + sinh Addendum (FR-12).''',
19: '''• FR-12: sinh Addendum CT-…-A1 sau khi trả phí gia hạn; deadline ký tại quầy 7 ngày.
• Amber banner trên Rental Detail + card Contract-signature trên Task Board + bell reminder hai bên.
• AWAITING_SIGNATURE đến khi staff đính kèm ảnh bản ký (nút hoàn tất vô hiệu đến khi capture tile có ảnh) → Signed, ảnh vào chuỗi.
• Quá hạn → reminder tăng cường, KHÔNG thu hồi ngày.
• Addendum không hold Unit; bị Voided (soạn nhầm / trả sớm) → soạn lại = bản ghi mới cùng Extension.''',
20: '''• FR-17: khách gửi yêu cầu trả kho từ Rental Detail kèm giải thích logic tất toán (hoàn / trừ / thu thêm).
• Submit → CHECKOUT_REQUESTED + Checkout Task sinh trên board.
• Request mới đè request cũ (cũ tự hủy); request DONE khi settlement hoàn tất.
• Ngày xin trả bị chặn trên biên Reservation kế tiếp (như Extension).''',
21: '''• FR-18: nhận Unit + key (checklist) → Inspection từng hạng mục (access card, padlock, cleanliness, structure — kết quả OK/MINOR/MAJOR) → khai Settlement Charge.
• Số tiền + reason BẮT BUỘC khi damage.
• Settlement preview: Deposit − charges = refund; charges vượt Deposit → khách trả phần chênh qua Payment Modal touchpoint Extra fee.
• KHÔNG đóng được khi có charge thiếu reason hoặc phần chênh chưa trả.
• Settlement receipt hiển thị hai bên vĩnh viễn ("Refund 63.500 ₫ after damage fee 40.000 ₫").
• Confirm đóng Rental → Unit → Preparing + Cleaning Task sinh ra.''',
22: '''• FR-19: hoàn tất Cleaning Task ngay trên card (không màn riêng).
• Hệ thống kiểm tra Turnover Buffer trước khi chuyển Unit → Available (hoặc Reserved nếu có Reservation kế).
• Buffer chưa đủ ngày → Unit giữ Preparing, card KHÔNG hoàn tất được.''',
23: '''• FR-23/24/35: khách tạo ticket (Unit + incident type LOST_ACCESS/DEVICE_ISSUE/SECURITY/CLEANLINESS/OTHER + mô tả) — không Unit hợp lệ không gửi được.
• Hệ thống route tới Staff trực ca theo Unit + shift.
• Staff Resolve kèm note (khách nhận kết quả plain words) hoặc Escalate KÈM NOTE BẮT BUỘC (không note → nút vô hiệu) → ticket vào Escalation Inbox của Facility Manager.
• Support List phía khách: danh sách + detail drawer nơi đọc resolution.''',
24: '''• FR-25: Manager đánh dấu severe → Unit chuyển Maintenance.
• Relocation = ĐỔI UNIT TRÊN RENTAL HIỆN TẠI (giữ mã RT-/Contract/Deposit, access code cấp mới, ActivityLog RELOCATION — không sinh Rental mới) + tasks bảo trì/dọn về kanban + khách nhận notification từng bước.
• Không severe → ticket về lại hàng staff kèm hướng dẫn.
• Sau mọi Severity Decision ticket quay IN_PROGRESS, chỉ Resolved khi staff hoàn tất + note.
• Ticket Resolved hiển thị trọn arc (sự cố → escalation → quyết định → bảo trì → di dời) trong một drawer.''',
25: '''• FR-26/20: drawer thao tác: sửa specs, Merge, Retire, set Maintenance, Fix status (reason BẮT BUỘC).
• Merge guard: không Rental/Reservation active hai bên + confirm "cannot be undone".
• Mọi chuyển trạng thái Unit qua guard + confirm — không silent status write; Merge/retire bị chặn → blocker nêu tên + link.
• Kết thúc Maintenance luôn qua Preparing + Cleaning Task theo buffer.
• Mọi write → Activity Log.''',
26: '''• FR-27: phân công staff × Zone × ca (MORNING/AFTERNOON/EVENING) × ngày.
• Conflict bị từ chối TRƯỚC KHI SAVE kèm collision cụ thể ("Minh is already on Morning, Zone B, Oct 12").
• Không lưu được phân công trùng người-ca-ngày.''',
27: '''• FR-28: audit trail append-only: timestamp, actor, entity, action, from → to, reason (nơi bắt buộc); lọc theo entity type.
• Read-only — không có UI xóa/sửa log (NFR-6: append-only ở tầng dữ liệu, chỉ INSERT).''',
28: '''• FR-37: tạo tài khoản cho mọi role với mật khẩu tạm; đổi role kèm Activity Log.
• Activate/deactivate/lock theo users.Status (0/1/2); KHÔNG delete — user chỉ đổi trạng thái; reset password = cấp mật khẩu tạm mới.
• Email trùng chặn inline; khoá chính mình bị chặn.
• Đăng nhập bằng tài khoản deactivate/lock → message chung không tiết lộ lý do.''',
29: '''• FR-38/39: xem event LOGIN/LOGIN_FAILED; filter theo user / thời gian / kết quả; mỗi dòng có actor, thời điểm, kết quả.
• Permission matrix Role × Permission (menu, hành động, phạm vi facility) hiển thị read-only đúng như server enforce — không chỉnh runtime.
• Pointer surface — hiển thị đúng nguồn, không tự tính lại.''',
30: '''• FR-29/30: sửa inline bảng rule (base rent theo Unit Type, phụ thu + trần, Deposit %, Turnover Buffer ngày — RuleType TURNOVER_BUFFER).
• Save chạy validation đầy đủ: giá trị vi phạm cắm cờ + banner nêu đúng luật vi phạm ("15% exceeds the 10% cap in Rental Policy v3") — policy xấu KHÔNG BAO GIỜ persist.
• Save thành công đóng dấu version + effective date; đổi buffer/deposit chỉ ảnh hưởng tính toán từ phiên bản effective.
• Block server hiển thị banner đầu form, không phải toast (NFR-8).''',
31: '''• FR-31: KPI revenue / deposits held / surcharges / utilization + charts.
• Reports: tab Revenue/Deposits/Surcharges/Occupancy + preset kỳ (this month, last month, quarter) + custom range + Export CSV (P2).
• Deposits held KHÔNG tự động tính là revenue.
• Mọi KPI bấm xuyên xuống bảng dữ liệu lọc đúng slice — không có số liệu nào không với tới được rows.''',
32: '''• FR-32: occupancy / revenue mix / unit-status KPI + charts + bảng nền theo cùng hợp đồng dashboard (KPI row → chart → table).
• Drill-down: KPI occupancy 87% → click → Unit Management lọc Rented.''',
33: '''• SM-1/2: chạy trọn 6 UJ end-to-end trên môi trường demo với seed chuẩn, không đứt gãy.
• Test phá các bước "ép": checkout charge thiếu reason, policy xấu lưu, card Done thiếu bước bật lại, access code cấp khi thiếu ảnh ký.
• Rà NFR: VND full precision + microcopy 3 phần (chuyện gì, hệ quả tiền/trạng thái, 1 bước kế) + a11y floor (label, focus, keyboard, status không chỉ màu).
• P2 đã cắt: Operations Monitor, notification day-grouping, Discount (FR-40), Waiver (FR-41) — cắt không phá flow.''',
}

NOTE = {
1: '''Monorepo theo Architecture Spine: backend/ Spring Boot 4.1.1 (Java 21), frontend/ Vite + React 19.3, contracts/openapi.yaml.
• Vite proxy /api → Boot :8080; FE phát triển chống MSW sinh từ contract.
• Cấm controller gọi thẳng repository — business logic nằm ở service (AD-3).''',
2: '''Flyway V1 = toàn bộ 22 entity theo model V3 (AD-6):
• Gộp RENTALS vào RESERVATIONS (status PENDING_PAYMENT/RESERVED/CHECKED_IN/CHECKOUT_REQUESTED/CLOSED/EXPIRED + AccessCode).
• Tách CONTRACT_ADDENDUMS (AWAITING_SIGNATURE/SIGNED/EXPIRED/VOIDED); CONTRACTS 1-1 RESERVATIONS (6 trạng thái + Superseded).
• POLICY_RULES.RuleType thêm TURNOVER_BUFFER/DISCOUNT/WAIVER_CAP; ActivityLogs.Reason NULL-able (NOT NULL chỉ cho status-change/charge).
• Cấm DDL tay, cấm hbm2ddl auto.
• US thuần BE, không có màn hình FE.''',
3: '''Nội dung seed: 5 role (kèm System Administrator); users demo Lan/Minh/Hằng/Tuấn/Nam; facility Tân Bình Depot + zones + units (S-3, M-2, M-5…); Rental Policy v3 (giá theo unit type, phụ thu + trần, Deposit 10%, Turnover Buffer ngày); reservations/payments lịch sử với mã chuẩn BK-1042, RT-0871, SR-0032, CT-1042, CT-1042-A1.
• US thuần BE, không có màn hình FE.''',
4: '''BCrypt hash mật khẩu at rest.
• FE: menu + role chip theo role; avatar menu (profile, logout).''',
5: '''FE: modal quên mật khẩu đặt trên màn Login.''',
6: '''FE: lưới card + chip "N units available · live"; filter active persist trong phiên, echo thành chip removable ở empty state.''',
7: '''FE: spec đầy đủ (kích thước, tầng, access, an ninh).''',
8: '''BE: interface PaymentGateway + MockPaymentGateway deterministic sau interface (AD-9) — FE không bao giờ tự quyết kết quả payment.''',
9: '—',
10: '''US thuần BE, không có màn hình FE.''',
11: '''US thuần BE, không có màn hình FE.''',
12: '—',
13: '—',
14: '''FE: 3px left bar màu theo loại card.''',
15: '—',
16: '—',
17: '—',
18: '—',
19: '—',
20: '—',
21: '—',
22: '—',
23: '—',
24: '—',
25: '''Bảng hiển thị: code mono, size, type, zone/floor, status badge, link rental/reservation active, last activity.''',
26: '''FE: lưới lịch tuần; slot xung đột highlight.''',
27: '—',
28: '''Bảng hiển thị: họ tên, email, phone, role, status, last login.''',
29: '—',
30: '—',
31: '—',
32: '—',
33: '''Chốt môi trường deploy (OQ-2: jar nhúng static/ vs docker-compose).
• Cả nhóm cùng thực hiện.''',
}

shutil.copyfile(SRC, BAK)
src = openpyxl.load_workbook(SRC, data_only=True)['Sprint Backlog']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sprint Backlog'

headers = ['US', 'Field/Screen', 'level', 'BE', 'FE', 'Sprint', 'From-To', 'Business Rules', 'Note']
widths  = [5, 26, 9, 9, 11, 9, 20, 72, 44]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.fill = PatternFill('solid', start_color='305496')
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

thin = Border(*[Side(style='thin', color='B0B0B0')]*4)
SPRINT_FILL = {'Sprint 1': 'DDEBF7', 'Sprint 2': 'E2EFDA', 'Sprint 3': 'FCE4D6'}
LEVEL_COLOR = {'complex': 'C00000', 'medium': 'BF8F00', 'simple': '548235'}

for r, row in enumerate(src.iter_rows(min_row=2, values_only=True), 2):
    us = row[0]
    for col, val in enumerate(row[:7], 1):
        ws.cell(row=r, column=col, value=val)
    ws.cell(row=r, column=8, value=BR[us])
    ws.cell(row=r, column=9, value=NOTE[us])
    for col in range(1, 10):
        c = ws.cell(row=r, column=col)
        c.border = thin
        c.alignment = Alignment(vertical='top', wrap_text=True,
                                horizontal='center' if col <= 7 else 'left')
    ws.cell(row=r, column=3).font = Font(bold=True, color=LEVEL_COLOR.get(row[2], '000000'))
    ws.cell(row=r, column=6).fill = PatternFill('solid', start_color=SPRINT_FILL.get(row[5], 'FFFFFF'))
    max_lines = 1
    for w, text in ((72, BR[us]), (44, NOTE[us])):
        for l in text.split('\n'):
            max_lines = max(max_lines, math.ceil(len(l) / (w - 2)))
    ws.row_dimensions[r].height = max(18, min(400, max_lines * 14.5 + 5))

ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:I{src.max_row}'
try:
    wb.save(SRC)
    print('OK: saved in place ->', SRC)
except PermissionError:
    alt = SRC.replace('.xlsx', '_v2.xlsx')
    wb.save(alt)
    print('LOCKED: file dang mo trong Excel -> da luu:', alt)
