# -*- coding: utf-8 -*-
"""Rewrite Sprint_Backlog.xlsx: Business Rules = short plain summary, Note = refs only."""
import shutil, math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"D:\FPT_Ky5\SWP\bmad swp\Sprint_Backlog.xlsx"
DET = r"D:\FPT_Ky5\SWP\bmad swp\.working\Sprint_Backlog_detailed.xlsx"

BR = {
1: 'Dựng khung dự án (monorepo BE + FE) và chốt hợp đồng API: mọi API lấy từ openapi.yaml, muốn đổi phải sửa contract trước. Thống nhất chuẩn lỗi + phân trang.',
2: 'Tạo toàn bộ CSDL (22 bảng) theo model V3 bằng Flyway. Tiền lưu DECIMAL cấm float; nhật ký chỉ ghi thêm, không sửa/xóa.',
3: 'Sinh dữ liệu demo: 5 quyền, user mẫu, kho Tân Bình + unit, bảng giá v3, dữ liệu lịch sử đủ chạy dashboard và demo 6 hành trình.',
4: 'Đăng nhập JWT: mỗi vai trò vào trang chủ riêng, sai thông tin báo lỗi chung, vào vùng khác vai trò bị 403. Phân quyền chặn ở server; ghi log mọi lần đăng nhập.',
5: 'Đăng ký công khai chỉ tạo Customer, email trùng báo ngay. Quên mật khẩu trả lời chung chung, chưa gửi email thật.',
6: 'Danh sách unit dạng card: lọc loại/kích thước/ngày, sort theo giá. Còn trống tính trên server (đặt chỗ + buffer dọn); unit không cho thuê không hiện.',
7: 'Trang chi tiết unit + bảng giá minh bạch: tiền thuê, từng phụ thu, cọc ghi rõ hoàn lại. Số tiền khớp mọi màn sau; format tiền VND chuẩn.',
8: 'Một modal thanh toán dùng chung: Card / MoMo / VNPay QR. Đang xử lý không đóng được; lỗi đồng nghĩa chưa trừ tiền, được thử lại hoặc đổi cách; thành công có receipt vĩnh viễn.',
9: 'Trước khi trả cọc khách duyệt lại toàn bộ dòng tiền. Bấm Reserve hệ thống kiểm lại unit còn trống — bị giành thì quay về lưới kèm gợi ý; cọc xong unit được giữ.',
10: 'Quá hạn nhận kho mà không check-in → Reservation hết hạn: mất cọc, unit mở lại, đóng contract draft, báo khách. Hết hạn tính trên server.',
11: 'Cọc thành công → tự động sinh hợp đồng draft, khóa phiên bản giá. Hợp đồng chỉ đọc; sai thì soạn lại, bản cũ vẫn xem được.',
12: 'Khách xem danh sách đặt chỗ/thuê của mình + mã check-in BK- đưa staff kiểm tại quầy kèm hướng dẫn. Số liệu lấy từ nguồn, không tự tính.',
13: 'Trang chi tiết một lần thuê: unit, ngày, lịch sử thanh toán, cọc, chuỗi hợp đồng + phụ lục. Mọi thao tác (check-in, gia hạn, trả kho, hỗ trợ) bắt đầu từ đây.',
14: 'Kanban cho staff: To do / In progress / Done với 5 loại thẻ (check-in, checkout, dọn, hỗ trợ, ký hợp đồng); di chuyển bằng kéo thả hoặc phím.',
15: 'Check-in tại quầy: xác nhận reservation đã cọc → thu 100% tiền thuê → in hợp đồng, chụp ảnh bản ký → mới được trao access code.',
16: 'Kéo thẻ sang Done khi chưa xong bước chốt (thanh toán, ảnh ký...) → thẻ bật lại và báo thiếu bước nào. Không có Done ảo.',
17: 'Thông báo: toast ngắn ~4s + chuông có badge số chưa đọc; feed theo vai trò, bấm vào tới thẳng đối tượng. Không gửi email.',
18: 'Gia hạn: khách chọn ngày trả mới, chặn nếu đụng đặt chỗ kế tiếp (báo rõ ngày giới hạn). Phí = tiền thuê kỳ thêm theo giá đang hiệu lực; trả xong cập nhật ngày + sinh phụ lục.',
19: 'Phụ lục gia hạn CT-…-A1: hẹn ký tại quầy 7 ngày, nhắc cả 2 bên; ký xong ảnh vào hồ sơ. Quá hạn chỉ nhắc mạnh hơn, không thu hồi ngày.',
20: 'Khách xin trả kho ngay trên trang rental → sinh thẻ Checkout cho staff; request mới thay request cũ.',
21: 'Staff nhận unit: checklist, kiểm tra hạng mục, kê khai phụ phí kèm lý do bắt buộc → tất toán (hoàn cọc hoặc thu thêm) → đóng rental, sinh thẻ dọn.',
22: 'Hoàn tất dọn ngay trên thẻ; unit chỉ mở cho thuê khi đủ ngày buffer theo policy.',
23: 'Khách gửi ticket gắn với unit; hệ thống chuyển cho staff trực ca. Staff xử lý kèm ghi chú hoặc chuyển lên quản lý (escalate phải có ghi chú).',
24: 'Quản lý đánh giá ticket: nghiêm trọng → unit chuyển bảo trì + chuyển khách sang unit khác ngay trên rental hiện tại (giữ hợp đồng, cọc); nhẹ → trả về staff.',
25: 'Quản lý unit: sửa thông tin, gộp, cho nghỉ, bảo trì — đều có xác nhận và chặn khi unit còn được giữ/thuê; mọi thao tác ghi log.',
26: 'Phân công staff theo khu vực × ca × ngày trên lịch tuần; trùng người-ca-ngày bị chặn trước khi lưu, báo cụ thể.',
27: 'Xem nhật ký hệ thống (ai làm gì, khi nào, vì sao) chỉ đọc, lọc theo đối tượng; không xóa/sửa được.',
28: 'SysAdmin quản lý user: tạo tài khoản mọi vai trò, đổi vai trò, khóa/mở, reset mật khẩu. Không xóa user — chỉ đổi trạng thái; không tự khóa chính mình.',
29: 'Xem lịch sử đăng nhập (thành/bại) + bảng quyền vai trò × chức năng đúng như server đang chặn, chỉ đọc.',
30: 'Sửa bảng giá quy định (giá cơ bản, phụ thu, % cọc, buffer ngày). Lưu luôn chạy validation — policy sai không bao giờ lưu được; lưu xong đóng dấu phiên bản + ngày hiệu lực.',
31: 'Dashboard kinh doanh: KPI doanh thu / cọc đang giữ / phụ thu / công suất + báo cáo theo kỳ, export CSV; bấm KPI xuyên xuống bảng chi tiết.',
32: 'Dashboard vận hành: công suất, cơ cấu doanh thu, trạng thái unit; drill-down (VD occupancy 87% → danh sách unit đang thuê).',
33: 'Chạy demo trọn 6 hành trình + test các ca bị ép (thiếu lý do charge, policy xấu, Done ảo...) + chốt môi trường deploy + rà NFR.',
}

NOTE = {
1: 'AD-2/3/8 · Spring Boot 4.1.1 (Java 21) + Vite React 19.3',
2: 'AD-6/7 · Flyway V1 · Thuần BE',
3: 'Mã demo chuẩn: BK-1042, RT-0871, SR-0032, CT-1042, CT-1042-A1 · Thuần BE',
4: 'FR-1, FR-38 · AD-5',
5: 'FR-2, FR-3 (P2)',
6: 'FR-4 · AD-4',
7: 'FR-6 · NFR-1',
8: 'FR-8/9 · AD-9 (gateway interface + mock)',
9: 'FR-5/7',
10: 'FR-36 · AD-4 · Thuần BE',
11: 'FR-10/13 · Thuần BE',
12: 'FR-35 · NFR-8',
13: 'FR-9/13/35',
14: 'FR-21 · NFR-2',
15: 'FR-11/14/15',
16: 'FR-22 · AD-4',
17: 'FR-33/34',
18: 'FR-16/12 · Trễ không có trạng thái OVERDUE — tính LATE_FEE lúc tất toán',
19: 'FR-12',
20: 'FR-17',
21: 'FR-18',
22: 'FR-19',
23: 'FR-23/24',
24: 'FR-25',
25: 'FR-26/20',
26: 'FR-27',
27: 'FR-28 · NFR-6',
28: 'FR-37',
29: 'FR-38/39',
30: 'FR-29/30 · NFR-8',
31: 'FR-31',
32: 'FR-32',
33: 'SM-1/2 · P2 đã cắt: Discount (FR-40), Waiver (FR-41), Operations Monitor',
}

shutil.copyfile(SRC, DET)
src = openpyxl.load_workbook(SRC, data_only=True)['Sprint Backlog']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sprint Backlog'

headers = ['US', 'Field/Screen', 'level', 'BE', 'FE', 'Sprint', 'From-To', 'Business Rules', 'Note']
widths  = [5, 26, 9, 9, 11, 9, 20, 80, 36]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.fill = PatternFill('solid', start_color='305496')
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = w

thin = Border(*[Side(style='thin', color='B0B0B0')]*4)
SPRINT_FILL = {'Sprint 1': 'DDEBF7', 'Sprint 2': 'E2EFDA', 'Sprint 3': 'FCE4D6'}
LEVEL_COLOR = {'complex': 'C00000', 'medium': 'BF8F00', 'simple': '548235'}

for r, row in enumerate(src.iter_rows(min_row=2, values_only=True), 2):
    us = row[0]
    for col, val in enumerate(row[:7], 1):
        ws.cell(row=r, column=col, value=val)
    ws.cell(row=r, column=8, value=BR[us])
    ws.cell(row=r, column=9, value=NOTE[us])
    max_lines = 1
    for w, text in ((80, BR[us]), (36, NOTE[us])):
        for l in text.split('\n'):
            max_lines = max(max_lines, math.ceil(len(l) / (w - 2)))
    ws.row_dimensions[r].height = max(18, min(120, max_lines * 14.5 + 5))
    for col in range(1, 10):
        c = ws.cell(row=r, column=col)
        c.border = thin
        c.alignment = Alignment(vertical='top', wrap_text=True,
                                horizontal='center' if col <= 7 else 'left')
    ws.cell(row=r, column=3).font = Font(bold=True, color=LEVEL_COLOR.get(row[2], '000000'))
    ws.cell(row=r, column=6).fill = PatternFill('solid', start_color=SPRINT_FILL.get(row[5], 'FFFFFF'))

ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:I{src.max_row}'
try:
    wb.save(SRC)
    print('OK: saved in place ->', SRC)
except PermissionError:
    alt = SRC.replace('.xlsx', '_v2.xlsx')
    wb.save(alt)
    print('LOCKED: file dang mo trong Excel -> da luu:', alt)
