# -*- coding: utf-8 -*-
"""Update Sprint_Backlog.xlsx: dời US-16 sang sprint 3, thêm US-34 PayOS sprint 2 (chốt 2026-09-23)."""
import copy
import openpyxl

PATH = r"D:\FPT_Ky5\SWP\StorageHub\docs\sprint\Sprint_Backlog.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Sprint Backlog"]

# --- 1) US-16 (row 17): dời sang sprint 3 ---
row16 = 17
assert ws.cell(row=row16, column=1).value == 16, "US-16 không nằm ở row 17"
ws.cell(row=row16, column=6).value = "Sprint 3"
ws.cell(row=row16, column=7).value = "20/10/2026 - 02/11/2026"
note16 = ws.cell(row=row16, column=9)
suffix = " · Dời sprint 3 (2026-09-23) nhường chỗ US-34 PayOS"
if suffix not in str(note16.value or ""):
    note16.value = (str(note16.value or "") + suffix).strip(" ·")

# --- 2) Thêm US-34 PayOS (append row 35), copy style từ row 18 (US-17, sprint 2) ---
style_row = 18
new_row = ws.max_row + 1
values = [
    34,
    "Payment — PayOS gateway thật",
    "complex",
    "Phúc",
    "Phú",
    "Sprint 2",
    "06/10/2026 - 19/10/2026",
    ("• PayOSGateway sau interface PaymentGateway — profile `payos` TẮT mặc định, demo chạy "
     "MockPaymentGateway (NFR-5).\n"
     "• Webhook POST /api/v1/payments/webhooks/payos = endpoint public thứ 4 — verify chữ ký "
     "HMAC-SHA256 bằng checksum key, idempotent theo mã đơn (giao dịch trùng = no-op).\n"
     "• PaymentService vẫn độc quyền state machine (AD-4) — gateway không bao giờ ghi DB.\n"
     "• FE modal render paymentUrl/QR khi server trả; không có → spinner mock như cũ; FE không "
     "tự quyết kết quả (AD-9).\n"
     "• Contract PR (webhook + paymentUrl) trước khi code (AD-2)."),
    ("FR-8/9 · AD-9 (sửa 2026-09-23) · Keys env: PAYOS_CLIENT_ID / PAYOS_API_KEY / PAYOS_CHECKSUM_KEY · "
     "Thêm 2026-09-23, đổi chỗ với US-16 (dời sprint 3) · Checkpoint 06/10: chưa có sandbox keys "
     "→ PayOS về stretch sprint 3, mock vẫn mặc định demo"),
    ("Tích hợp cổng thanh toán thật PayOS sau seam PaymentGateway sẵn có: tạo payment request nhận "
     "link/QR, nhận webhook xác nhận (verify chữ ký, idempotent), giữ mock làm mặc định demo. "
     "Đăng ký PayOS sandbox lấy keys ngay tuần đầu sprint 1 (việc duyệt có thể chậm); thử tunnel "
     "(ngrok/cloudflared) với endpoint test."),
]
for col, val in enumerate(values, start=1):
    src = ws.cell(row=style_row, column=col)
    dst = ws.cell(row=new_row, column=col)
    dst.value = val
    dst.font = copy.copy(src.font)
    dst.alignment = copy.copy(src.alignment)
    dst.border = copy.copy(src.border)
    dst.fill = copy.copy(src.fill)
    dst.number_format = src.number_format

wb.save(PATH)
print("saved OK")

# --- verify ---
wb2 = openpyxl.load_workbook(PATH)
ws2 = wb2["Sprint Backlog"]
print("US-16 →", ws2.cell(row=17, column=6).value, "|", ws2.cell(row=17, column=7).value)
r = ws2.max_row
print("new row", r, ":", ws2.cell(row=r, column=1).value, "|", ws2.cell(row=r, column=2).value,
      "|", ws2.cell(row=r, column=3).value, "| BE", ws2.cell(row=r, column=4).value,
      "| FE", ws2.cell(row=r, column=5).value, "|", ws2.cell(row=r, column=6).value)
