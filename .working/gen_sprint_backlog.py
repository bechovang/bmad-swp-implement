# -*- coding: utf-8 -*-
"""Sinh Sprint_Backlog.xlsx cho StorageHub từ planning artifacts (PRD 41 FR + Architecture Spine)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.load_workbook("Sprint_Backlog.xlsx")
ws = wb["Sprint Backlog"]

# Xóa dòng mẫu (giữ header dòng 1)
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

S1 = ("Sprint 1", "22/09/2026 - 05/10/2026")
S2 = ("Sprint 2", "06/10/2026 - 19/10/2026")
S3 = ("Sprint 3", "20/10/2026 - 02/11/2026")

# (Field/Screen, level, BE, FE, sprint, Business Rules)
ROWS = [
    # ---------- SPRINT 1: Nền tảng + Identity + Discovery & Booking ----------
    ("Project Setup + API Contract", "medium", "An", "Phú", S1,
     "Dựng monorepo theo Architecture Spine: backend/ Spring Boot 4.1.1 (Java 21), frontend/ Vite + React 19.3, contracts/openapi.yaml là chân lý API duy nhất (AD-2) — đóng băng đầu sprint, muốn đổi thì PR vào contract trước. Lỗi + list theo envelope chung (AD-8: machine code + message + field errors; pagination {items, page, pageSize, total}, mặc định 25 rows). Vite proxy /api → Boot :8080; FE phát triển chống MSW sinh từ contract. Cấm controller gọi thẳng repository — business logic nằm ở service (AD-3)."),
    ("DB Schema — Flyway V1 (model V3)", "complex", "Phúc", "—", S1,
     "Flyway V1 = toàn bộ 22 entity model V3 (AD-6): gộp RENTALS vào RESERVATIONS (status PENDING_PAYMENT/RESERVED/CHECKED_IN/CHECKOUT_REQUESTED/CLOSED/EXPIRED + AccessCode), tách CONTRACT_ADDENDUMS (AWAITING_SIGNATURE/SIGNED/EXPIRED/VOIDED), CONTRACTS 1-1 RESERVATIONS (6 trạng thái + Superseded), POLICY_RULES.RuleType thêm TURNOVER_BUFFER/DISCOUNT/WAIVER_CAP, ActivityLogs.Reason NULL-able (NOT NULL chỉ cho status-change/charge). Tiền = DECIMAL(15,0) — cấm float/double (AD-7); thời gian lưu UTC. Cấm DDL tay, cấm hbm2ddl auto. ActivityLog append-only: chỉ INSERT."),
    ("Seed demo data", "medium", "Huy", "—", S1,
     "Seed 5 role (kèm System Administrator); users demo Lan/Minh/Hằng/Tuấn/Nam; facility Tân Bình Depot + zones + units (S-3, M-2, M-5…); Rental Policy v3 (giá theo unit type, phụ thu + trần, Deposit 10%, Turnover Buffer ngày); dữ liệu reservations/payments lịch sử đủ chạy dashboard + 6 UJ (mã chuẩn BK-1042, RT-0871, SR-0032, CT-1042, CT-1042-A1)."),
    ("Login + Adaptive Shell 5 role", "complex", "An", "Phú", S1,
     "FR-1: JWT Bearer, role trong claim, TTL 24h, không refresh token; landing theo role (Customer → Browse Units; Staff → Task Board; Facility Manager → Facility Overview; Business Ops → Business Overview; SysAdmin → User Management). Sai thông tin → message chung không tiết lộ field nào sai; truy cập URL chéo role → 403. Menu + role chip theo role, avatar menu (profile, logout). BCrypt hash at rest; permission matrix enforce server-side từng endpoint (AD-5). Mỗi lần đăng nhập (thành/bại) ghi ActivityLog LOGIN/LOGIN_FAILED (FR-38)."),
    ("Register + Forgot password", "simple", "Huy", "Tuấn Anh", S1,
     "FR-2: form công khai chỉ tạo CUSTOMER — không có trường chọn role; email trùng chặn inline; thành công tự đăng nhập → Browse Units. FR-3 (P2): modal quên mật khẩu trên Login, phản hồi generic \"If an account exists…\", không gửi email thật."),
    ("Browse Units + availability chính xác", "complex", "Phúc", "Tuấn Anh", S1,
     "FR-4: filter loại / kích thước / ngày bắt đầu / thời hạn — chỉ re-query khi bấm Search (không per keystroke); lưới card + chip \"N units available · live\"; sort mặc định giá thấp → cao. Availability = Reservation + Turnover Buffer tính tại thời điểm query, on-read idempotent trong service (AD-4) — FE hiển thị nguyên văn không tính lại. Unit Rented/Maintenance/Retired không xuất hiện; card buffer \"Available {date} · cleaning buffer\" vẫn đặt được với ngày đó; filter active persist trong phiên, echo thành chip removable ở empty state."),
    ("Unit Detail + bảng giá minh bạch", "medium", "Phúc", "Tuấn Anh", S1,
     "FR-6: spec đầy đủ (kích thước, tầng, access, an ninh) + bảng giá: rent × duration + từng dòng phụ thu từ Rental Policy active + Deposit đánh dấu refundable. Mọi dòng tiền khớp tuyệt đối Booking Summary và Payment Modal. Tiền VND full precision, format 1.150.000 ₫, tabular numerals (NFR-1)."),
    ("Payment Modal + mock gateway", "complex", "An", "Phú", S1,
     "FR-8/9: một modal dùng chung 4 touchpoint (Deposit, 100% rent, Extension fee, Extra fee); 3 phương thức Card (validation inline) / MoMo (phone + OTP) / VNPay QR (QR + đếm ngược ~5 phút). Processing ~1.5–2s — modal không đóng/điều hướng được; fail = \"No money was taken\" + Retry + Switch method, KHÔNG đổi trạng thái gì anywhere; QR hết hạn → về method select; sau 2 fail gợi ý đổi phương thức. BE: interface PaymentGateway + MockPaymentGateway deterministic sau interface (AD-9) — FE không bao giờ tự quyết kết quả payment. Mỗi payment thành công sinh đúng 1 receipt xem vĩnh viễn."),
    ("Booking Summary + Reserve re-check", "complex", "Phúc", "Tuấn Anh", S1,
     "FR-5/7: khách duyệt lại toàn bộ dòng sẽ bị truy đòi trước khi trả cọc — không có dòng tiền nào xuất hiện ở payment mà vắng ở đây; ghi chú contract được auto-draft từ đúng các điều kiện này và ký tại check-in. Bấm Reserve → hệ thống re-check availability: unit bị chiếm giữa chừng → quay về lưới + toast gợi ý unit tương tự, không bao giờ booking trên ngày invalid. Sinh Reservation PENDING_PAYMENT ngay tại Booking Summary; cọc thành công → RESERVED (unit bị giữ từ đây)."),
    ("Reservation lifecycle + EXPIRED no-show", "complex", "An", "—", S1,
     "FR-36: hết ngày nhận kho mà chưa check-in → EXPIRED: mất cọc (receipt \"Deposit forfeited — no-show\"), Unit → Available, Contract Draft/Printed → Closed, ghi Activity Log + notification khách. EXPIRED được tính on-read idempotent (AD-4) — API luôn trả state đã suy diễn xong, FE không tự suy từ ngày. Không có cancel trong v1 (PRD §6)."),
    ("Contract auto-draft", "complex", "Phúc", "—", S1,
     "FR-10: sinh Contract Draft ngay khi Deposit thành công; mã CT-xxxx; khóa phiên bản Rental Policy; read-only ở mọi nơi — không UI nào cho phép sửa, không ai soạn contract tay. Sai sót bản Draft sửa bằng re-draft: bản cũ chuyển Superseded, vẫn đọc được trong chuỗi (FR-13)."),
    ("My Rentals + Check-in Pass", "medium", "Huy", "Phú", S1,
     "FR-35: My Rentals — danh sách Reservations + Rentals + History dạng card + status badge + nút action kế tiếp; là pointer surface: mọi con số/trạng thái khớp nguồn, không tự tính lại. Check-in Pass — mã đặt chỗ BK- hiển thị lớn (mono) để staff validate tại quầy + hướng dẫn plain-English (mang ID ký hợp đồng, trả 100% tiền thuê tại quầy). Empty state có đúng 1 CTA trỏ flow đổ dữ liệu (NFR-8)."),
    ("Rental Detail + contract chain", "complex", "Huy", "Phú", S1,
     "FR-9/13/35: một rental = unit, dates, payment history + receipts vĩnh viễn, deposit status (held/settled), contract chain: contract gốc + addenda (mã CT-, badge trạng thái, ảnh bản ký); bản Signed xem mãi; Superseded vẫn đọc được; addendum chưa ký có amber banner nêu deadline tại quầy. Là anchor của mọi action: Check-in, Extend, Checkout Request, New Support — không flow nào bắt đầu từ system menu."),

    # ---------- SPRINT 2: Check-in → Checkout trọn vòng đời + Kanban + Notifications ----------
    ("Task Board — kanban 5 loại card", "complex", "Phúc", "Tuấn Anh", S2,
     "FR-21: cột To do / In progress / Done; 5 loại card Check-in / Checkout / Cleaning / Support / Contract-signature (3px left bar màu theo loại); tab lọc theo loại; số lượng từng cột ở header; kéo thả HOẶC keyboard Move (drag không phải con đường duy nhất — NFR-2); card Done render muted. Card sinh theo ca trực (route theo Unit + shift)."),
    ("Check-in Task — ritual hợp đồng", "complex", "An", "Tuấn Anh", S2,
     "FR-11/14/15: nhập/mã reservation → validate tồn tại + Deposit đã trả, không hợp lệ chặn kèm lý do cụ thể (chưa cọc không check-in được). Thu 100% rent qua Payment Modal — Deposit và rent full là 2 giá trị tách bạch. Contract ritual: preview read-only → Print → chụp ảnh bản ký → Attach (ghi CONTRACT_SIGNED vào Activity Log); nút trao Access Code VÔ HIỆU đến khi có ảnh bản ký. Hoàn tất → Rental CHECKED_IN, Unit RENTED, task Done."),
    ("Snap-back chống Done ảo", "medium", "Phúc", "Tuấn Anh", S2,
     "FR-22: kéo Check-in/Checkout/Contract-signature vào Done khi thiếu closing step (payment/settlement/ảnh bản ký) → card bật lại + toast nêu đúng bước thiếu. Không có task Done mà nghiệp vụ chưa khép — enforce phía service (AD-4), FE hiển thị kết quả."),
    ("Toast + Bell + Notification feed", "medium", "Huy", "Phú", S2,
     "FR-33/34: toast tức thời ~4s, hover pause, ≤1 action link; bell badge đếm unread; Notification Center feed theo role (khách: booking/payment/support; staff: task assignment; manager: escalation + status writes), unread trước, deep link tới đối tượng, mark-all-read, unread persist giữa phiên. Sự kiện tiền phát CẢ toast lẫn bell. Không email, không bao giờ."),
    ("Extension + conflict boundary", "complex", "An", "Phú", S2,
     "FR-16: khách chọn ngày checkout mới → hệ thống kiểm Reservation kế tiếp; vùng xung đột đánh dấu trên picker; submit bị chặn kèm biên rõ (\"Latest new checkout: Oct 18\"). Chỉ mở TRƯỚC EndDate — qua EndDate giữ CHECKED_IN, mỗi ngày trễ tính phụ thu LATE_FEE tại Settlement (không có trạng thái OVERDUE). Phí gia hạn = tiền thuê kỳ thêm theo Rental Policy active (không dùng công thức cọc mới − cọc cũ); Deposit giữ nguyên booking gốc. Thanh toán thành công → EndDate flip NGAY + sinh Addendum (FR-12)."),
    ("Addendum giấy + Contract-signature card", "complex", "An", "Phú", S2,
     "FR-12: sinh Addendum CT-…-A1 sau khi trả phí gia hạn; deadline ký tại quầy 7 ngày; amber banner trên Rental Detail + card Contract-signature trên Task Board + bell reminder hai bên. AWAITING_SIGNATURE đến khi staff đính kèm ảnh bản ký (nút hoàn tất vô hiệu đến khi capture tile có ảnh) → Signed, ảnh vào chuỗi. Quá hạn → reminder tăng cường, KHÔNG thu hồi ngày. Addendum không hold Unit; bị Voided (soạn nhầm / trả sớm) → soạn lại = bản ghi mới cùng Extension."),
    ("Checkout Request", "medium", "Huy", "Phú", S2,
     "FR-17: khách gửi yêu cầu trả kho từ Rental Detail kèm giải thích logic tất toán (hoàn / trừ / thu thêm) → CHECKOUT_REQUESTED + Checkout Task sinh trên board. Request mới đè request cũ (cũ tự hủy); request DONE khi settlement hoàn tất; ngày xin trả bị chặn trên biên Reservation kế tiếp (như Extension)."),
    ("Checkout Task — climax tất toán", "complex", "Phúc", "Tuấn Anh", S2,
     "FR-18: nhận Unit + key (checklist) → Inspection từng hạng mục (access card, padlock, cleanliness, structure — kết quả OK/MINOR/MAJOR) → khai Settlement Charge: số tiền + reason BẮT BUỘC khi damage → settlement preview (Deposit − charges = refund; charges vượt Deposit → khách trả phần chênh qua Payment Modal touchpoint Extra fee) → confirm đóng Rental. KHÔNG đóng được khi có charge thiếu reason hoặc phần chênh chưa trả; settlement receipt hiển thị hai bên vĩnh viễn (\"Refund 63.500 ₫ after damage fee 40.000 ₫\"); Unit → Preparing + Cleaning Task sinh ra."),
    ("Cleaning + Turnover buffer", "medium", "Huy", "Tuấn Anh", S2,
     "FR-19: hoàn tất Cleaning Task ngay trên card (không màn riêng); hệ thống kiểm tra Turnover Buffer trước khi chuyển Unit → Available (hoặc Reserved nếu có Reservation kế). Buffer chưa đủ ngày → Unit giữ Preparing, card KHÔNG hoàn tất được."),

    # ---------- SPRINT 3: Quản trị + Dashboard + Support/Escalation + Demo ----------
    ("Support Ticket — tạo, route, resolve/escalate", "complex", "Phúc", "Phú", S3,
     "FR-23/24/35: khách tạo ticket (Unit + incident type LOST_ACCESS/DEVICE_ISSUE/SECURITY/CLEANLINESS/OTHER + mô tả) — không Unit hợp lệ không gửi được; hệ thống route tới Staff trực ca theo Unit + shift. Staff Resolve kèm note (khách nhận kết quả plain words) hoặc Escalate KÈM NOTE BẮT BUỘC (không note → nút vô hiệu) → ticket vào Escalation Inbox của Facility Manager. Support List phía khách: danh sách + detail drawer nơi đọc resolution."),
    ("Severity Decision + Relocation", "complex", "An", "Tuấn Anh", S3,
     "FR-25: Manager đánh dấu severe → Unit chuyển Maintenance + Relocation = ĐỔI UNIT TRÊN RENTAL HIỆN TẠI (giữ mã RT-/Contract/Deposit, access code cấp mới, ActivityLog RELOCATION — không sinh Rental mới) + tasks bảo trì/dọn về kanban + khách nhận notification từng bước. Không severe → ticket về lại hàng staff kèm hướng dẫn. Sau mọi Severity Decision ticket quay IN_PROGRESS, chỉ Resolved khi staff hoàn tất + note; ticket Resolved hiển thị trọn arc (sự cố → escalation → quyết định → bảo trì → di dời) trong một drawer."),
    ("Unit Management + edit drawer + guards", "complex", "Phúc", "Tuấn Anh", S3,
     "FR-26/20: bảng Unit (code mono, size, type, zone/floor, status badge, link rental/reservation active, last activity) + drawer: sửa specs, Merge (guard: không Rental/Reservation active hai bên + confirm \"cannot be undone\"), Retire (guard + confirm), set Maintenance, Fix status (reason BẮT BUỘC). Mọi chuyển trạng thái Unit qua guard + confirm — không silent status write; kết thúc Maintenance luôn qua Preparing + Cleaning Task theo buffer. Merge/retire bị chặn → blocker nêu tên + link. Mọi write → Activity Log."),
    ("Staff & Shifts + conflict detection", "medium", "Huy", "Tuấn Anh", S3,
     "FR-27: phân công staff × Zone × ca (MORNING/AFTERNOON/EVENING) × ngày; conflict bị từ chối TRƯỚC KHI SAVE kèm collision cụ thể (\"Minh is already on Morning, Zone B, Oct 12\") + slot xung đột highlight; lưới lịch tuần. Không lưu được phân công trùng người-ca-ngày."),
    ("Activity Log viewer", "simple", "Huy", "Phú", S3,
     "FR-28: audit trail append-only: timestamp, actor, entity, action, from → to, reason (nơi bắt buộc); lọc theo entity type; read-only — không có UI xóa/sửa log (NFR-6: append-only ở tầng dữ liệu, chỉ INSERT)."),
    ("User Management — SYS-01", "medium", "An", "Phú", S3,
     "FR-37: bảng user (họ tên, email, phone, role, status, last login) + search/filter theo role & status; TẠO tài khoản cho mọi role với mật khẩu tạm; đổi role kèm Activity Log; activate/deactivate/lock theo users.Status (0/1/2); KHÔNG có delete — user chỉ đổi trạng thái; reset password = cấp mật khẩu tạm mới. Email trùng chặn inline; khoá chính mình bị chặn; đăng nhập bằng tài khoản deactivate/lock → message chung không tiết lộ lý do."),
    ("Login & Activity History — SYS-02 + Permission matrix", "simple", "Huy", "Phú", S3,
     "FR-38/39: xem event LOGIN/LOGIN_FAILED từ Activity Log; filter theo user / thời gian / kết quả; mỗi dòng có actor, thời điểm, kết quả; pointer surface — hiển thị đúng nguồn, không tự tính lại. Permission matrix Role × Permission (menu, hành động, phạm vi facility) hiển thị read-only đúng như server enforce — không chỉnh runtime."),
    ("Policy Management — validate trước khi lưu", "complex", "Phúc", "Phú", S3,
     "FR-29/30: sửa inline bảng rule (base rent theo Unit Type, phụ thu + trần, Deposit %, Turnover Buffer ngày — RuleType TURNOVER_BUFFER); Save chạy validation đầy đủ: giá trị vi phạm cắm cờ + banner nêu đúng luật vi phạm (\"15% exceeds the 10% cap in Rental Policy v3\") — policy xấu KHÔNG BAO GIỜ persist; Save thành công đóng dấu version + effective date. Đổi buffer/deposit chỉ ảnh hưởng tính toán từ phiên bản effective. Block server hiển thị banner đầu form, không phải toast (NFR-8)."),
    ("Business Overview + Reports + KPI drill-down", "complex", "An", "Tuấn Anh", S3,
     "FR-31: KPI revenue / deposits held / surcharges / utilization + charts; Reports: tab Revenue/Deposits/Surcharges/Occupancy + preset kỳ (this month, last month, quarter) + custom range + Export CSV (P2). Deposits held KHÔNG tự động tính là revenue; mọi KPI bấm xuyên xuống bảng dữ liệu lọc đúng slice — không có số liệu nào không với tới được rows."),
    ("Facility Overview + drill-down", "medium", "Huy", "Tuấn Anh", S3,
     "FR-32: occupancy / revenue mix / unit-status KPI + charts + bảng nền theo cùng hợp đồng dashboard (KPI row → chart → table). Drill-down: KPI occupancy 87% → click → Unit Management lọc Rented."),
    ("E2E Demo 6 UJ + deploy + hardening", "complex", "Cả nhóm", "Cả nhóm", S3,
     "SM-1/2: chạy trọn 6 UJ end-to-end trên môi trường demo với seed chuẩn, không đứt gãy; test phá các bước \"ép\": checkout charge thiếu reason, policy xấu lưu, card Done thiếu bước bật lại, access code cấp khi thiếu ảnh ký. Chốt môi trường deploy (OQ-2: jar nhúng static/ vs docker-compose); rà NFR: VND full precision + microcopy 3 phần (chuyện gì, hệ quả tiền/trạng thái, 1 bước kế) + a11y floor (label, focus, keyboard, status không chỉ màu). P2 đã cắt: Operations Monitor, notification day-grouping, Discount (FR-40), Waiver (FR-41) — cắt không phá flow."),
]

LEVEL_FILL = {
    "simple": PatternFill("solid", fgColor="E2EFDA"),
    "medium": PatternFill("solid", fgColor="FFF2CC"),
    "complex": PatternFill("solid", fgColor="FCE4D6"),
}
BASE_FILL = PatternFill("solid", fgColor="F9FBFD")
thin = Side(style="thin", color="B0B7C3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def est_height(screen, rules):
    lines_b = max(1, -(-len(screen) // 34))   # col B width 32
    lines_h = max(1, -(-len(rules) // 88))    # col H width 85
    return max(20, max(lines_b, lines_h) * 13.2 + 6)

for i, (screen, level, be, fe, (sprint, dates), rules) in enumerate(ROWS, start=2):
    ws.cell(row=i, column=1, value=i - 1)
    ws.cell(row=i, column=2, value=screen)
    ws.cell(row=i, column=3, value=level)
    ws.cell(row=i, column=4, value=be)
    ws.cell(row=i, column=5, value=fe)
    ws.cell(row=i, column=6, value=sprint)
    ws.cell(row=i, column=7, value=dates)
    ws.cell(row=i, column=8, value=rules)
    for col in range(1, 9):
        c = ws.cell(row=i, column=col)
        c.font = Font(name="Segoe UI", size=10, bold=(col in (2,)))
        c.border = BORDER
        c.alignment = LEFT if col in (2, 8) else CENTER
        c.fill = LEVEL_FILL[level] if col == 3 else BASE_FILL
    ws.row_dimensions[i].height = est_height(screen, rules)

ws.column_dimensions["E"].width = 14
ws.column_dimensions["F"].width = 12

wb.save("Sprint_Backlog.xlsx")
print(f"OK — wrote {len(ROWS)} user stories: Sprint1={sum(1 for r in ROWS if r[4] is S1)}, "
      f"Sprint2={sum(1 for r in ROWS if r[4] is S2)}, Sprint3={sum(1 for r in ROWS if r[4] is S3)}")
