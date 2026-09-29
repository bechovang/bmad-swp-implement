# -*- coding: utf-8 -*-
"""Rewrite backlog: Mo ta (process) + Business Rules (atomic, testable) + Note (refs)."""
import shutil, math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"D:\FPT_Ky5\SWP\bmad swp\Sprint_Backlog.xlsx"
BAK = r"D:\FPT_Ky5\SWP\bmad swp\.working\Sprint_Backlog_simplified.xlsx"

RULES = {
1: '''• openapi.yaml là nguồn chân lý duy nhất của API — muốn đổi API phải sửa contract trước khi code.
• Mọi lỗi trả về theo envelope chung: machine code + message + field errors.
• Mọi danh sách trả về {items, page, pageSize, total}, mặc định 25 dòng/trang.''',
2: '''• Tiền tệ lưu DECIMAL(15,0) — cấm float/double.
• Mọi mốc thời gian lưu UTC.
• ActivityLog chỉ được INSERT — không UPDATE, không DELETE.
• Thay đổi schema chỉ qua migration Flyway — cấm DDL tay và hbm2ddl auto.''',
3: '''• Seed phải đủ 5 role (kèm System Administrator) và user demo cho từng role.
• Seed phải đủ dữ liệu lịch sử để dashboard và 6 UJ chạy trọn không đứt gãy.''',
4: '''• Phiên đăng nhập hiệu lực 24h (JWT) — không cấp refresh token.
• Mỗi vai trò vào đúng trang chủ riêng; truy cập URL của vai trò khác → 403.
• Đăng nhập sai chỉ trả message chung — không tiết lộ field nào sai.
• Quyền truy cập kiểm tra ở server theo permission matrix, áp cho mọi endpoint.
• Mọi lần đăng nhập (thành/bại) phải ghi ActivityLog.
• Mật khẩu chỉ lưu dạng BCrypt hash.''',
5: '''• Form đăng ký công khai chỉ tạo CUSTOMER — không có trường chọn role.
• Email đã tồn tại → chặn, báo inline ngay tại form.
• Đăng ký thành công → tự đăng nhập vào Browse Units.
• Quên mật khẩu luôn trả phản hồi giống nhau dù email có tồn tại hay không.
• Không gửi email thật trong v1 (P2).''',
6: '''• Unit Rented/Maintenance/Retired không xuất hiện trong kết quả.
• Availability = Reservation + Turnover Buffer, tính tại thời điểm query trên server — FE hiển thị nguyên văn, không tự tính lại.
• Unit đang trong buffer dọn hiển thị ngày sẵn sàng và vẫn đặt được với ngày đó.
• Chỉ làm mới danh sách khi bấm Search — không query theo từng phím gõ.
• Sắp xếp mặc định: giá thấp → cao.''',
7: '''• Bảng giá phải đủ: rent × duration + từng dòng phụ thu theo Rental Policy đang hiệu lực.
• Deposit hiển thị kèm nhãn refundable (hoàn lại).
• Mọi dòng tiền phải khớp tuyệt đối với Booking Summary và Payment Modal.
• Tiền hiển thị VND full precision (1.150.000 ₫).''',
8: '''• Một modal dùng chung cho 4 điểm thanh toán: Deposit, 100% rent, Extension fee, Extra fee.
• Hỗ trợ 3 phương thức: Card / MoMo (OTP) / VNPay QR (hết hạn sau ~5 phút).
• Đang xử lý không được đóng modal hay điều hướng rời trang.
• Thất bại = chưa trừ tiền, KHÔNG thay đổi trạng thái nào — khách được Retry hoặc đổi phương thức.
• Sau 2 lần thất bại phải gợi ý đổi phương thức.
• Mỗi payment thành công sinh đúng 1 receipt, xem được vĩnh viễn.
• Kết quả thanh toán chỉ do server (gateway) quyết định — FE không tự quyết.''',
9: '''• Booking Summary phải liệt kê 100% dòng tiền khách sẽ bị truy đòi — không dòng nào xuất hiện ở Payment mà vắng ở đây.
• Hợp đồng draft chỉ sinh từ đúng các điều kiện đã duyệt tại đây.
• Bấm Reserve phải re-check availability — unit bị chiếm giữa chừng thì không được đặt, quay về lưới + gợi ý unit tương tự.
• Không bao giờ booking trên ngày không còn hiệu lực.
• Reserve thành công → Reservation = PENDING_PAYMENT.
• Deposit thành công → Reservation = RESERVED, unit bị giữ từ đây.''',
10: '''• Quá ngày nhận kho mà chưa check-in → Reservation = EXPIRED.
• EXPIRED → khách mất cọc, receipt ghi "Deposit forfeited — no-show".
• EXPIRED → Unit = Available; Contract Draft/Printed = Closed.
• EXPIRED phải ghi ActivityLog + notification cho khách.
• EXPIRED do server suy diễn on-read — API luôn trả state cuối, FE không tự suy từ ngày.
• v1 không có cancel đặt chỗ.''',
11: '''• Contract Draft sinh tự động ngay khi Deposit thành công — không có đường sinh tay.
• Mã contract dạng CT-xxxx.
• Mỗi contract khóa phiên bản Rental Policy tại thời điểm sinh.
• Contract luôn read-only — không UI nào cho phép sửa.
• Sửa draft = re-draft: bản cũ chuyển Superseded, vẫn đọc được trong chuỗi.''',
12: '''• My Rentals hiển thị đủ 3 nhóm: Reservations / Rentals / History, mỗi card có status badge + nút action kế tiếp.
• Mọi số liệu/trạng thái lấy nguyên văn từ nguồn — FE không tự tính lại.
• Mã đặt chỗ BK- hiển thị lớn (mono) để staff đối chiếu tại quầy.
• Empty state chỉ có đúng 1 CTA trỏ về flow đổ dữ liệu.''',
13: '''• Mỗi rental hiển thị: unit, dates, payment history, receipts vĩnh viễn, deposit status (held/settled).
• Chuỗi contract: gốc + mọi addendum; bản Signed xem mãi; Superseded vẫn đọc được.
• Addendum chưa ký phải có amber banner nêu deadline ký tại quầy.
• Mọi action (Check-in, Extend, Checkout Request, New Support) chỉ bắt đầu từ đây — không flow nào bắt đầu từ system menu.''',
14: '''• Board đúng 3 cột: To do / In progress / Done.
• Thẻ đúng 5 loại: Check-in / Checkout / Cleaning / Support / Contract-signature.
• Thẻ sinh tự động theo ca trực (route theo Unit + shift).
• Di chuyển thẻ bằng kéo thả HOẶC bàn phím — kéo thả không phải con đường duy nhất.
• Thẻ Done hiển thị muted; số lượng từng cột hiển thị ở header.''',
15: '''• Chỉ được check-in khi Reservation tồn tại VÀ Deposit đã thanh toán — vi phạm chặn kèm lý do cụ thể.
• Deposit và 100% rent là 2 khoản thanh toán tách biệt.
• Access code chỉ được trao sau khi có ảnh bản ký hợp đồng.
• Đính kèm ảnh bản ký phải ghi CONTRACT_SIGNED vào ActivityLog.
• Hoàn tất → Rental = CHECKED_IN, Unit = RENTED, task = Done.''',
16: '''• Thẻ chỉ được Done khi bước chốt đã xong (thanh toán / settlement / ảnh bản ký).
• Kéo sang Done khi thiếu bước chốt → thẻ bật lại + toast nêu đúng bước thiếu.
• Rule enforce ở service — FE chỉ hiển thị kết quả.''',
17: '''• Toast tồn tại ~4s, hover pause, tối đa 1 action link.
• Bell badge đếm số chưa đọc; unread persist giữa các phiên.
• Feed theo vai trò: khách (booking/payment/support), staff (task assignment), manager (escalation + status writes).
• Sự kiện liên quan tiền phải phát cả toast lẫn bell.
• Feed sắp unread trước, deep link tới đối tượng.
• Không gửi email — không bao giờ.''',
18: '''• Ngày checkout mới không được đụng Reservation kế tiếp — submit bị chặn kèm biên rõ ("Latest new checkout: Oct 18").
• Extension chỉ mở TRƯỚC EndDate hiện tại.
• Quá EndDate: giữ CHECKED_IN, không có trạng thái OVERDUE — mỗi ngày trễ tính LATE_FEE lúc settlement.
• Phí gia hạn = tiền thuê kỳ thêm theo Rental Policy đang hiệu lực (không dùng hiệu số cọc).
• Deposit giữ nguyên booking gốc.
• Thanh toán thành công → EndDate cập nhật NGAY + sinh Addendum.''',
19: '''• Addendum (CT-…-A1) sinh sau khi trả phí gia hạn.
• Deadline ký tại quầy: 7 ngày.
• Trạng thái AWAITING_SIGNATURE cho đến khi có ảnh bản ký → SIGNED.
• Nút hoàn tất ký vô hiệu đến khi capture tile có ảnh.
• Quá deadline → reminder tăng cường, KHÔNG thu hồi ngày.
• Addendum không hold Unit.
• Addendum Voided (soạn nhầm/trả sớm) → soạn lại thành bản ghi mới của cùng Extension.''',
20: '''• Checkout request gửi từ Rental Detail, kèm giải thích logic tất toán (hoàn / trừ / thu thêm).
• Gửi thành công → Rental = CHECKOUT_REQUESTED + sinh Checkout Task.
• Request mới thay request cũ — request cũ tự hủy.
• Request chỉ DONE khi settlement hoàn tất.
• Ngày xin trả bị chặn trên biên Reservation kế tiếp (như Extension).''',
21: '''• Nhận Unit + key theo checklist; Inspection từng hạng mục (access card, padlock, cleanliness, structure) với kết quả OK/MINOR/MAJOR.
• Charge damage bắt buộc có số tiền + reason.
• Settlement = Deposit − charges = refund; charges vượt Deposit → khách trả phần chênh (touchpoint Extra fee).
• Không đóng rental khi còn charge thiếu reason hoặc phần chênh chưa trả.
• Settlement receipt hiển thị 2 bên vĩnh viễn.
• Đóng rental → Unit = Preparing + sinh Cleaning Task.''',
22: '''• Cleaning hoàn tất ngay trên thẻ — không màn hình riêng.
• Unit chỉ rời Preparing khi Turnover Buffer đã đủ ngày theo policy.
• Hết buffer → Unit = Available, hoặc Reserved nếu có Reservation kế.''',
23: '''• Ticket bắt buộc: Unit hợp lệ + incident type (LOST_ACCESS/DEVICE_ISSUE/SECURITY/CLEANLINESS/OTHER) + mô tả.
• Không có Unit hợp lệ → không gửi được.
• Ticket route tự động tới staff trực ca theo Unit + shift.
• Resolve phải kèm note — khách nhận kết quả plain words.
• Escalate bắt buộc note (không note → nút vô hiệu) → vào Escalation Inbox của Facility Manager.''',
24: '''• Đánh dấu severe → Unit = Maintenance.
• Relocation = đổi unit trên Rental hiện tại: giữ mã RT-/Contract/Deposit — không sinh Rental mới.
• Relocation: access code mới + ActivityLog RELOCATION + tasks bảo trì/dọn về kanban + notification khách từng bước.
• Không severe → ticket về lại staff kèm hướng dẫn.
• Sau Severity Decision ticket luôn quay IN_PROGRESS — chỉ Resolved khi staff hoàn tất + note.''',
25: '''• Merge chỉ khi không có Rental/Reservation active ở cả 2 unit + confirm "cannot be undone".
• Mọi chuyển trạng thái unit qua guard + confirm — không silent status write.
• Fix status bắt buộc reason.
• Bị chặn → blocker nêu tên + link.
• Kết thúc Maintenance luôn qua Preparing + Cleaning Task theo buffer.
• Mọi thao tác ghi Activity Log.''',
26: '''• Phân công theo Staff × Zone × Ca (MORNING/AFTERNOON/EVENING) × ngày.
• Không lưu phân công trùng người-ca-ngày — chặn trước khi SAVE, báo collision cụ thể.''',
27: '''• Mỗi dòng log: timestamp, actor, entity, action, from → to, reason (nơi bắt buộc).
• Viewer read-only — không tồn tại UI xóa/sửa.
• Lọc được theo entity type.''',
28: '''• Tạo được tài khoản cho mọi role, kèm mật khẩu tạm.
• Đổi role phải ghi ActivityLog.
• Trạng thái user chỉ 0/1/2 (deactivate/active/lock) — KHÔNG xóa user.
• Reset password = cấp mật khẩu tạm mới.
• Email trùng → chặn inline.
• Không được tự khóa chính mình.
• Đăng nhập bằng tài khoản deactivate/lock → message chung, không tiết lộ lý do.''',
29: '''• Hiển thị LOGIN/LOGIN_FAILED với actor, thời điểm, kết quả; filter theo user/thời gian/kết quả.
• Permission matrix Role × Permission hiển thị đúng như server enforce.
• Matrix read-only — không chỉnh runtime; không tự tính lại nguồn.''',
30: '''• Chỉnh inline: base rent theo Unit Type, phụ thu + trần, Deposit %, Turnover Buffer ngày.
• Save luôn chạy validation — giá trị vi phạm bị cắm cờ + banner nêu đúng luật ("15% exceeds the 10% cap in Rental Policy v3").
• Policy không hợp lệ không bao giờ được persist.
• Save thành công → đóng dấu version + effective date.
• Đổi buffer/deposit chỉ ảnh hưởng tính toán từ phiên bản effective.
• Block server hiển thị banner đầu form — không dùng toast.''',
31: '''• KPI: revenue / deposits held / surcharges / utilization.
• Deposits held không tự động tính là revenue.
• Reports 4 tab + preset kỳ + custom range + export CSV.
• Mọi KPI phải xuyên xuống bảng dữ liệu lọc đúng slice — không số liệu nào không với tới được.''',
32: '''• KPI occupancy / revenue mix / unit-status.
• Dashboard theo cùng hợp đồng: KPI row → chart → table.
• Drill-down: bấm KPI occupancy → Unit Management lọc Rented.''',
33: '''• 6 UJ phải chạy trọn trên môi trường demo với seed chuẩn, không đứt gãy.
• Test bắt buộc các ca ép: charge thiếu reason, policy xấu lưu, card Done thiếu bước, trao access code khi thiếu ảnh ký.
• NFR phải rà: VND full precision, microcopy 3 phần, a11y floor (label, focus, keyboard, status không chỉ màu).
• P2 đã cắt: Operations Monitor, notification day-grouping, Discount (FR-40), Waiver (FR-41) — cắt không phá flow.''',
}

shutil.copyfile(SRC, BAK)
src = openpyxl.load_workbook(SRC, data_only=True)['Sprint Backlog']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Sprint Backlog'

headers = ['US', 'Field/Screen', 'level', 'BE', 'FE', 'Sprint', 'From-To', 'Mô tả', 'Business Rules', 'Note']
widths  = [5, 25, 9, 9, 11, 9, 20, 52, 78, 34]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.fill = PatternFill('solid', start_color='305496')
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = w

thin = Border(*[Side(style='thin', color='B0B0B0')]*4)
SPRINT_FILL = {'Sprint 1': 'DDEBF7', 'Sprint 2': 'E2EFDA', 'Sprint 3': 'FCE4D6'}
LEVEL_COLOR = {'complex': 'C00000', 'medium': 'BF8F00', 'simple': '548235'}

for r, row in enumerate(src.iter_rows(min_row=2, values_only=True), 2):
    us, desc, note = row[0], row[7], row[8]
    for col, val in enumerate(row[:7], 1):
        ws.cell(row=r, column=col, value=val)
    ws.cell(row=r, column=8, value=desc)        # Mô tả (quy trình ngắn, từ bản trước)
    ws.cell(row=r, column=9, value=RULES[us])   # Business Rules (atomic)
    ws.cell(row=r, column=10, value=note)       # Note (refs)
    max_lines = 1
    for w, text in ((52, desc), (78, RULES[us])):
        for l in str(text).split('\n'):
            max_lines = max(max_lines, math.ceil(len(l) / (w - 2)))
    ws.row_dimensions[r].height = max(18, min(220, max_lines * 14.5 + 5))
    for col in range(1, 11):
        c = ws.cell(row=r, column=col)
        c.border = thin
        c.alignment = Alignment(vertical='top', wrap_text=True,
                                horizontal='center' if col <= 7 else 'left')
    ws.cell(row=r, column=3).font = Font(bold=True, color=LEVEL_COLOR.get(row[2], '000000'))
    ws.cell(row=r, column=6).fill = PatternFill('solid', start_color=SPRINT_FILL.get(row[5], 'FFFFFF'))

ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:J{src.max_row}'
try:
    wb.save(SRC)
    print('OK: saved in place ->', SRC)
except PermissionError:
    alt = SRC.replace('.xlsx', '_v2.xlsx')
    wb.save(alt)
    print('LOCKED: file dang mo trong Excel -> da luu:', alt)
