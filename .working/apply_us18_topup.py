# -*- coding: utf-8 -*-
"""US 18 (row 19): áp dụng công thức top-up cọc (quyết định 2026-09-22)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

PATH = "sprintlog Swp chỉnh sửa.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Sprint Backlog"]

ws["H19"] = (
    "• Ngày checkout mới không được đụng Reservation kế tiếp — submit bị chặn kèm biên rõ "
    "(\"Latest new checkout: Oct 18\").\n"
    "• Extension chỉ mở TRƯỚC EndDate hiện tại.\n"
    "• Phí gia hạn = tiền thuê kỳ thêm theo Rental Policy đang hiệu lực + top-up cọc "
    "= max(0, Deposit% × tổng tiền thuê hợp đồng sau gia hạn − cọc đang giữ).\n"
    "• Thanh toán thành công → EndDate cập nhật NGAY, cọc đang giữ nâng lên mức mới, "
    "sinh Addendum (ghi số cọc mới).\n"
    "• Settlement dùng số cọc hiện giữ — công thức refund không đổi."
)
ws["I19"] = (
    "FR-16/12 · Trễ không có OVERDUE — tính LATE_FEE lúc tất toán · "
    "Top-up cọc chốt 2026-09-22 (demo: 690.000 + 69.000 = 759.000 ₫, cọc → 172.500 ₫)"
)
ws["J19"] = (
    "Gia hạn: khách chọn ngày trả mới, chặn nếu đụng đặt chỗ kế tiếp (báo rõ ngày giới hạn). "
    "Phí = tiền thuê kỳ thêm + top-up cọc về mức mới theo giá đang hiệu lực; "
    "trả xong cập nhật ngày + cọc + sinh phụ lục."
)

wb.save(PATH)
print("OK — US 18 (H19/I19/J19) updated với công thức top-up.")
