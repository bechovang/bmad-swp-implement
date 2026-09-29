# -*- coding: utf-8 -*-
"""Áp dụng các quyết định chốt 2026-09-22 vào sprintlog Swp chỉnh sửa.xlsx
- US 14 (row 15): bỏ "Và do Manager giao" → mọi thẻ do hệ thống sinh từ sự kiện nghiệp vụ
- US 19 (row 20): Voided bỏ trigger "soạn nhầm", giữ "trả kho sớm"; phục hồi bullet "không hold Unit"
- US 25 (row 26): Merge → P2, không build
(US 18 row 19: chưa đụng — chờ chốt phương án cọc)
"""
import shutil, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

PATH = "sprintlog Swp chỉnh sửa.xlsx"
shutil.copy(PATH, ".working/sprintlog_backup_2026-09-22.xlsx")

wb = openpyxl.load_workbook(PATH)
ws = wb["Sprint Backlog"]

# --- US 14 (row 15): Task Board ---
ws["H15"] = (
    "• Board đúng 3 cột: To do / In progress / Done.\n"
    "• Thẻ đúng 5 loại: Check-in / Checkout / Cleaning / Support / Contract-signature.\n"
    "• Mọi thẻ do hệ thống sinh từ sự kiện nghiệp vụ, route theo Unit + shift "
    "(kể cả task bảo trì/dọn phát sinh từ Severity Decision của Manager) — không có đường giao tay.\n"
    "• Mặc định: hiện thẻ trong ngày/ca hiện tại."
)

# --- US 19 (row 20): Addendum ---
ws["H20"] = (
    "• Addendum (CT-…-A1) sinh tự động sau khi trả phí gia hạn — không có đường sinh tay "
    "nên không tồn tại ca \"soạn nhầm\" (khác hợp đồng gốc có re-draft).\n"
    "• Addendum không hold Unit — ngày mới hiệu lực ngay từ lúc trả phí, phụ lục chỉ là hồ sơ giấy.\n"
    "• Deadline ký tại quầy: 7 ngày.\n"
    "• Trạng thái AWAITING_SIGNATURE cho đến khi có ảnh bản ký → SIGNED.\n"
    "• Nút hoàn tất ký vô hiệu đến khi capture tile có ảnh.\n"
    "• Voided chỉ còn 1 đường: khách trả kho sớm trước khi ký → staff đóng hồ sơ, không soạn lại."
)
ws["I20"] = "FR-12 · Chốt 2026-09-22: bỏ trigger \"soạn nhầm\", giữ Voided cho trả kho sớm"

# --- US 25 (row 26): Unit Management — Merge → P2 ---
ws["H26"] = (
    "• Merge 2 unit → P2, KHÔNG build (chốt 2026-09-22 — cắt không phá flow; "
    "beat demo UJ-4 thay bằng Retire S-2, vẫn demo đủ guard + blocker + confirm).\n"
    "• Mọi chuyển trạng thái unit qua guard + confirm — không silent status write.\n"
    "• Fix status bắt buộc reason.\n"
    "• Bị chặn → blocker nêu tên + link.\n"
    "• Kết thúc Maintenance luôn qua Preparing + Cleaning Task theo buffer.\n"
    "• Mọi thao tác ghi Activity Log."
)
ws["J26"] = (
    "Quản lý unit: sửa thông tin, cho nghỉ, bảo trì, sửa trạng thái — đều có xác nhận và "
    "chặn khi unit còn được giữ/thuê; mọi thao tác ghi log. (Merge 2 unit: P2 — không build.)"
)
ws["I26"] = "FR-26/20 · Merge → P2 2026-09-22"

wb.save(PATH)
print("OK — updated H15, H20/I20, H26/I26/J26. Backup: .working/sprintlog_backup_2026-09-22.xlsx")
